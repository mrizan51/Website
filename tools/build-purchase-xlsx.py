#!/usr/bin/env python3
"""Builds PEI-Purchase-Statement.xlsx — the purchase statement cleaned and made
consistent, ready to go live on Google Sheets with purchase-statement/Code.gs.

    python tools/build-purchase-xlsx.py [source.xlsx] [out.xlsx] [--dump DIR]

The cleaning is the same, rule for rule and word for word in the log, as
cleanData_() in Code.gs, and the party names are read from that file, so the
two can never disagree about a name. --dump writes the source and the expected
result as JSON for tools/test-purchase-gs.js, which runs Code.gs against them.
"""
import sys, os, re, json, datetime as dt
from collections import OrderedDict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GS = os.path.join(ROOT, 'purchase-statement', 'Code.gs')
args = [a for i, a in enumerate(sys.argv[1:], start=1)
        if not a.startswith('--') and sys.argv[i - 1] != '--dump']
SRC = args[0] if args else os.path.join(ROOT, 'purchase-statement', 'Purchase-original.xlsx')
OUT = args[1] if len(args) > 1 else os.path.join(ROOT, 'purchase-statement', 'PEI-Purchase-Statement.xlsx')
DUMP = sys.argv[sys.argv.index('--dump') + 1] if '--dump' in sys.argv else None

# ---------------------------------------------------------------- from Code.gs
gs = open(GS, encoding='utf-8').read()
def const(name, open_, close):
    m = re.search('const ' + name + ' = (' + re.escape(open_) + '.*?' + re.escape(close) + ');', gs, re.S)
    if not m:
        sys.exit(f'{name} not found in {GS}')
    return json.loads(m.group(1))
PARTY_NAMES = const('PARTY_NAMES', '{', '}')
CONFIRM = const('CONFIRM', '[', ']')
num = lambda name: int(re.search(rf'const {name} = (\d+);', gs).group(1))
FIRST, ROWS, HEAD_ROW = num('FIRST'), num('ROWS'), num('HEAD_ROW')
TITLE = re.search(r"const TITLE = '([^']*)';", gs).group(1)
LOG_HEAD = json.loads(re.search(r"const LOG_HEAD = (\[[^\]]*\]);", gs).group(1).replace("'", '"'))
LAST = FIRST + ROWS - 1

def key(s):
    return ' '.join(str(s).split()).upper()

INDEX = {}
for k, v in PARTY_NAMES.items():
    INDEX[key(k)] = v
    INDEX[key(v)] = v
CONFIRM_KEYS = {key(k) for k in CONFIRM}

def canonical(s):
    t = ' '.join(str(s).split())
    return INDEX.get(key(t), t)

def lot_text(n):
    return f'PUR-{int(round(n)):03d}'

def lot_number(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(round(v))
    m = re.fullmatch(r'PUR-?\s*0*(\d+)', str(v).strip(), re.I)
    return int(m.group(1)) if m else None

blank = lambda v: v is None or v == ''
money = lambda x: f'{x:,.2f}'
day = lambda d: d.strftime('%d/%m/%Y')

# ---------------------------------------------------------------- read
src = load_workbook(SRC)
ws = src['Purchase']
def cellv(r, c):
    v = ws.cell(row=r, column=c).value
    return '' if v is None else v

last = FIRST - 1
for r in range(FIRST, ws.max_row + 1):
    if any(not blank(cellv(r, c)) for c in range(2, 10)):
        last = r
source = {r: [cellv(r, c) for c in range(1, 12)] for r in range(1, ws.max_row + 1)}
rows = [[cellv(r, c) for c in range(2, 10)] for r in range(FIRST, last + 1)]   # B..I
B, C, D, E, F, G, H, I = range(8)

# ---------------------------------------------------------------- clean
now = dt.datetime.now().replace(microsecond=0)
log, numbered = [], 0
for i, row in enumerate(rows):
    r = FIRST + i
    if isinstance(row[I], str) and row[I] != row[I].strip():
        t = row[I].strip()
        log.append([now, r, '', 'Notes', row[I], t, 'Spaces trimmed' if t else 'Only spaces — cleared'])
        row[I] = t
    if isinstance(row[E], str) and row[E].strip():
        m = re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*box(?:es)?\s*', row[E], re.I)
        if m:
            note = f'{m.group(1)} boxes'
            log.append([now, r, '', 'Weight (kg)', row[E], '', f'Boxes, not kg — moved to Notes as "{note}"'])
            row[I] = f'{note}; {row[I]}' if row[I] else note
            row[E] = ''
        elif re.fullmatch(r'\s*\d+(?:\.\d+)?\s*', row[E]):
            log.append([now, r, '', 'Weight (kg)', row[E], float(row[E]), 'Text made a number'])
            row[E] = float(row[E])
    if isinstance(row[D], str) and row[D].strip():
        c = canonical(row[D])
        if c != row[D]:
            log.append([now, r, '', 'Party', row[D], c,
                        'Same supplier by timing and price, not spelling — please confirm'
                        if key(row[D]) in CONFIRM_KEYS else 'Same supplier, another spelling'])
            row[D] = c
    has = any(not blank(row[k]) for k in (C, D, E, F, G, H, I))
    n = None if blank(row[B]) else lot_number(row[B])
    if not has:
        if not blank(row[B]):
            log.append([now, r, '', 'Lot No.', lot_text(n) if n is not None else str(row[B]), '',
                        'Number on an empty row — removed; the next purchase is numbered automatically'])
            row[B] = ''
    elif n is not None and row[B] != lot_text(n):
        row[B] = lot_text(n)
        numbered += 1

isdate = lambda v: isinstance(v, dt.datetime)
isnum = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
same1, same2 = OrderedDict(), OrderedDict()
for i, row in enumerate(rows):
    if not isdate(row[C]) or not isnum(row[F]):
        continue
    same1.setdefault((row[C], key(row[D]), row[F]), []).append(i)
    if isnum(row[E]):
        same2.setdefault((row[C], row[E], row[F]), []).append(i)
other = lambda group, i: next((j for j in group if j != i), None)
for i, row in enumerate(rows):
    r, lot = FIRST + i, row[B]
    check = lambda note: log.append([now, r, lot, 'Check', '', '', note])
    if isdate(row[H]) and isdate(row[C]) and row[H] < row[C]:
        check(f'Payment date {day(row[H])} is before the purchase date {day(row[C])}')
    if isnum(row[E]) and row[E] > 0 and isnum(row[F]):
        rate = row[F] / row[E]
        if rate < 20 or rate > 2000:
            check(f'Weight looks wrong for the amount: ₹{money(rate)} per kg')
    if isnum(row[F]) and isnum(row[G]) and abs(row[F] - row[G]) > 0.005:
        check(f'Amount and Swipe differ by ₹{money(abs(row[F] - row[G]))}')
    if isdate(row[C]) and isnum(row[F]):
        j1 = other(same1[(row[C], key(row[D]), row[F])], i)
        j2 = other(same2[(row[C], row[E], row[F])], i) if isnum(row[E]) else None
        if j1 is not None:
            check(f'Same date, party and amount as {rows[j1][B]}')
        elif j2 is not None:
            check(f'Same date, weight and amount as {rows[j2][B]} ({rows[j2][D]})')
    if isinstance(row[I], str) and re.search(r'Rs\.\s*/-', row[I]):
        check(f'Note has no amount: "{row[I]}"')
if numbered:
    log.append([now, '', '', 'Lot No.', 'numbers', 'PUR- text',
                f'{numbered} lot numbers stored as text with the PUR- prefix'])
log.append([now, '', '', 'Month', 'typed labels', 'formula', 'Month shown on every row, worked out from the Date'])

parties = sorted({key(canonical(row[D])): canonical(row[D]) for row in rows
                  if isinstance(row[D], str) and row[D].strip()}.values(), key=str.upper)

# ---------------------------------------------------------------- style
FONT, SIZE = 'Arial', 10
INK, MUTED, TITLE_C, HEAD, RULE = '0B2C3A', '5E7C8B', '065E7A', '065E7A', 'D5E3EA'
RED_FILL, RED_INK, GREEN_FILL, GREEN_INK, AMBER = 'F4C7C3', 'A50E0E', 'B7E1CD', '0D652D', 'FCE8B2'
# (heading, width px, alignment, number format) — the COLUMNS table in Code.gs
COLUMNS = [('Month', 70, 'center', None), ('Lot No.', 95, 'center', '@'),
           ('Date', 95, 'center', 'DD/MM/YYYY'), ('Party', 230, 'left', '@'),
           ('Weight (kg)', 105, 'right', '#,##0.00'), ('Amount (₹)', 125, 'right', '[$₹]#,##0.00'),
           ('Swipe (₹)', 125, 'right', '[$₹]#,##0.00'), ('Payment', 100, 'center', 'DD/MM/YYYY'),
           ('Notes', 260, 'left', '@')]
gs_cols = re.findall(r"\['([^']+)', (\d+), '(\w+)', (null|'[^']*')\]", gs)
if [(h, int(w), a) for h, w, a, _ in gs_cols] != [(h, w, a) for h, w, a, _ in COLUMNS]:
    sys.exit('COLUMNS differ from Code.gs')
px = lambda p: round((p - 5) / 7, 2)       # Google pixels → Excel character widths
fill = lambda c: PatternFill('solid', fgColor=c)
rule = Side(style='thin', color=RULE)

wb = Workbook()
sh = wb.active
sh.title = 'Purchase'
sh.sheet_view.showGridLines = False
sh.freeze_panes = f'A{FIRST}'

sh['A1'] = TITLE
sh['A1'].font = Font(name=FONT, size=12, bold=True, color=TITLE_C)
sh['A1'].alignment = Alignment(vertical='center')
sh.row_dimensions[1].height = 22.5
for c, (head, w, align, fmt) in enumerate(COLUMNS, start=1):
    h = sh.cell(row=HEAD_ROW, column=c, value=head)
    h.font = Font(name=FONT, size=SIZE, bold=True, color='FFFFFF')
    h.fill = fill(HEAD)
    h.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    sh.column_dimensions[get_column_letter(c)].width = px(w)
sh.row_dimensions[HEAD_ROW].height = 25.5

body_font = Font(name=FONT, size=SIZE, color=INK)
month_font = Font(name=FONT, size=SIZE, color=MUTED)
aligns = [Alignment(horizontal=a, vertical='center') for _, _, a, _ in COLUMNS]
bottom = Border(bottom=rule)
for r in range(FIRST, LAST + 1):
    vals = rows[r - FIRST] if r <= last else [''] * 8
    m = sh.cell(row=r, column=1, value=f'=IF(C{r}="","",UPPER(TEXT(C{r},"mmm")))')
    m.font, m.alignment, m.border = month_font, aligns[0], bottom
    for c in range(2, 10):
        v = vals[c - 2]
        cell = sh.cell(row=r, column=c, value=None if v == '' else v)
        cell.font, cell.alignment, cell.border = body_font, aligns[c - 1], bottom
        cell.number_format = COLUMNS[c - 1][3]

# Amount greater → Swipe red; Swipe greater → Amount green; a likely double entry → amber Lot No.
rng = lambda col: f'{col}{FIRST}:{col}{LAST}'
sh.conditional_formatting.add(rng('G'), FormulaRule(
    formula=[f'AND(ISNUMBER($F{FIRST}),$F{FIRST}>N($G{FIRST}))'],
    fill=fill(RED_FILL), font=Font(name=FONT, bold=True, color=RED_INK)))
sh.conditional_formatting.add(rng('F'), FormulaRule(
    formula=[f'AND(ISNUMBER($G{FIRST}),$G{FIRST}>N($F{FIRST}))'],
    fill=fill(GREEN_FILL), font=Font(name=FONT, bold=True, color=GREEN_INK)))
sh.conditional_formatting.add(rng('B'), FormulaRule(
    formula=[f'AND($C{FIRST}<>"",$D{FIRST}<>"",ISNUMBER($F{FIRST}),'
             f'COUNTIFS($C${FIRST}:$C${LAST},$C{FIRST},$D${FIRST}:$D${LAST},$D{FIRST},'
             f'$F${FIRST}:$F${LAST},$F{FIRST})>1)'],
    fill=fill(AMBER)))

def dv(ref, **k):
    v = DataValidation(allow_blank=True, showErrorMessage=True, showInputMessage=True, **k)
    sh.add_data_validation(v)
    v.add(ref)
dv(rng('B'), type='custom',
   formula1=f'AND(LEFT(B{FIRST},4)="PUR-",ISNUMBER(--MID(B{FIRST},5,9)),LEN(B{FIRST})>=7,'
            f'COUNTIF($B${FIRST}:$B${LAST},B{FIRST})=1)',
   prompt='Leave blank: the next PUR- number fills in by itself. If typed, it must be PUR- and a number not used before, e.g. PUR-746.',
   error='Lot No. must be PUR- and a number not used before, e.g. PUR-746.')
dv(rng('C'), type='date', operator='between', formula1='DATE(2020,1,1)', formula2='DATE(2040,12,31)',
   prompt='The purchase date, e.g. 05/10/2026.', error='Enter a date.')
dv(rng('D'), type='list', formula1=f'Parties!$A$2:$A${len(parties) + 200}',
   prompt='Pick the party from the list. A new party is added at the foot of the Parties sheet first.',
   error='Pick a party from the list.')
for col, msg in (('E', 'Kilograms, as a number. A count of boxes goes in Notes.'),
                 ('F', 'Rupees, as a number.'), ('G', 'Rupees, as a number.')):
    dv(rng(col), type='decimal', operator='greaterThanOrEqual', formula1='0', prompt=msg,
       error='Enter a number, 0 or more.')
dv(rng('H'), type='custom', formula1=f'AND(ISNUMBER(H{FIRST}),OR($C{FIRST}="",H{FIRST}>=$C{FIRST}))',
   prompt='The payment date, on or after the purchase date.', error='Payment date must be a date on or after the purchase date.')
sh.print_title_rows = f'{HEAD_ROW}:{HEAD_ROW}'
sh.page_setup.orientation = 'landscape'
sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
sh.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- Parties
ps = wb.create_sheet('Parties')
ps.sheet_view.showGridLines = False
ps.freeze_panes = 'A2'
ps['A1'] = 'Party'
ps['A1'].font = Font(name=FONT, size=SIZE, bold=True, color='FFFFFF')
ps['A1'].fill = fill(HEAD)
ps['B1'] = 'One name per supplier. Add a new one in the first empty cell; it then appears in the Party list.'
ps['B1'].font = Font(name=FONT, size=SIZE, italic=True, color=MUTED)
ps.column_dimensions['A'].width = px(260)
ps.column_dimensions['B'].width = px(560)
for i in range(max(len(parties) + 200, 300)):
    c = ps.cell(row=2 + i, column=1, value=parties[i] if i < len(parties) else None)
    c.font, c.border, c.number_format = body_font, bottom, '@'
ps.conditional_formatting.add(f'A2:A{len(parties) + 200}', FormulaRule(
    formula=[f'AND($A2<>"",COUNTIF($A$2:$A${len(parties) + 200},$A2)>1)'],
    fill=fill(RED_FILL), font=Font(name=FONT, color=RED_INK)))

# ---------------------------------------------------------------- Cleanup log
ls = wb.create_sheet('Cleanup log')
ls.sheet_view.showGridLines = False
ls.freeze_panes = 'A2'
for c, h in enumerate(LOG_HEAD, start=1):
    x = ls.cell(row=1, column=c, value=h)
    x.font = Font(name=FONT, size=SIZE, bold=True, color='FFFFFF')
    x.fill = fill(HEAD)
for c, w in enumerate((130, 50, 90, 90, 210, 190, 520), start=1):
    ls.column_dimensions[get_column_letter(c)].width = px(w)
for i, line in enumerate(log, start=2):
    for c, v in enumerate(line, start=1):
        x = ls.cell(row=i, column=c, value=None if v == '' else v)
        x.font = Font(name=FONT, size=SIZE, color=INK, bold=(c == 4 and v == 'Check'))
        if c == 1:
            x.number_format = 'DD/MM/YYYY HH:MM'
wb.active = 0
wb.save(OUT)

changes = [l for l in log if l[3] != 'Check']
print(f'wrote {OUT}: {len(rows)} rows ({sum(1 for r in rows if any(not blank(v) for v in r[1:]))} purchases), '
      f'{len(parties)} parties (from {len({source[r][3] for r in source if r >= FIRST and isinstance(source[r][3], str) and source[r][3].strip()})} spellings), '
      f'{len(changes)} changes, {len(log) - len(changes)} checks')

# ---------------------------------------------------------------- for the Code.gs test
if DUMP:
    os.makedirs(DUMP, exist_ok=True)
    enc = lambda v: {'$date': v.strftime('%Y-%m-%d')} if isinstance(v, dt.datetime) else v
    json.dump({'maxRows': ws.max_row, 'maxCols': ws.max_column,
               'cells': {f'{r},{c + 1}': enc(v) for r, vals in source.items()
                         for c, v in enumerate(vals) if not blank(v)}},
              open(os.path.join(DUMP, 'purchase-source.json'), 'w'), ensure_ascii=False)
    json.dump({'first': FIRST, 'last': last, 'rows': [[enc(v) for v in row] for row in rows],
               'parties': parties, 'log': [[enc(v) for v in l[1:]] for l in log]},
              open(os.path.join(DUMP, 'purchase-expected.json'), 'w'), ensure_ascii=False)
    print(f'dumped test data to {DUMP}')
