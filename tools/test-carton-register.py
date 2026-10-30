#!/usr/bin/env python3
"""Checks the carton register against an independent calculation.

    python tools/test-carton-register.py live.xlsx WORKDIR

Builds the register from the live stock register with the count date set to
01/10/2026, so a month of real production is deducted; then types a working
scenario — opening counts, reorder levels, deliveries in and out of the date
window, a damaged-carton adjustment, a mistyped carton, a carton shared by two
brands, a product packed with no carton set — recalculates in LibreOffice, and
compares every carton figure with the same arithmetic done here in Python.
"""
import sys, os, math, subprocess, json, datetime as dt
from openpyxl import load_workbook

SRC, WORK = sys.argv[1], sys.argv[2]
HERE = os.path.dirname(os.path.abspath(__file__))
RECALC = next(p for p in ('/mnt/skills/public/xlsx/scripts/recalc.py',
                          os.path.expanduser('~/.claude/skills/xlsx/scripts/recalc.py')) if os.path.exists(p))
os.makedirs(WORK, exist_ok=True)
built, case = os.path.join(WORK, 'carton-test.xlsx'), os.path.join(WORK, 'carton-case.xlsx')
COUNT = dt.datetime(2026, 10, 1)
subprocess.run([sys.executable, os.path.join(HERE, 'add-carton-register.py'), SRC, built,
                '--count-date', COUNT.strftime('%Y-%m-%d')], check=True)

# ---------------------------------------------------------------- the scenario, typed as a person would
wb = load_workbook(built)
cs, ce, pk, st, de = wb['Carton Stock'], wb['Carton Entry'], wb['Packing'], wb['Stock'], wb['Daily Entry']
STMT = st['E3'].value
cartons = {cs.cell(row=r, column=2).value: r for r in range(8, 48) if cs.cell(row=r, column=2).value}
opening = {'Blue Dolphin carton': 1200, 'Primus Japan carton': 800, 'Primus EU carton': 300,
           'Tuna carton': 150, 'China carton': 5}
reorder = {'Blue Dolphin carton': 2000, 'Primus Japan carton': 100, 'China carton': 50}
for n, v in opening.items(): cs.cell(row=cartons[n], column=3).value = v
for n, v in reorder.items(): cs.cell(row=cartons[n], column=10).value = v
deliveries = [  # date, carton, received, adjustment
    (dt.datetime(2026, 9, 30), 'Blue Dolphin carton', 999, None),    # before the count: not counted
    (dt.datetime(2026, 10, 5), 'Blue Dolphin carton', 500, None),
    (dt.datetime(2026, 10, 12), 'Primus Japan carton', 250, -4),     # 4 damaged on arrival
    (STMT, 'Primus EU carton', 100, None),                           # today
    (STMT, 'Tuna carton', None, 3),                                  # found in a recount
    (STMT + dt.timedelta(days=1), 'Primus EU carton', 77, None),     # after the statement date: not yet
    (dt.datetime(2026, 10, 20), 'Blue Dolphn carton', 60, None),     # mistyped: flagged, not counted
]
for i, (d, n, rec, adj) in enumerate(deliveries):
    r = 7 + i
    ce.cell(row=r, column=1).value, ce.cell(row=r, column=2).value = d, n
    ce.cell(row=r, column=3).value, ce.cell(row=r, column=4).value = rec, adj
# AMF packs into Blue Dolphin cartons; tuna 20 kg a carton; mackerel left without a size
for r in range(8, 102):
    b = st.cell(row=r, column=1).value
    if b == 'AMF Brand': pk.cell(row=r, column=4).value = 'Blue Dolphin carton'
    if b == 'Tuna': pk.cell(row=r, column=5).value = 20
wb.save(case)
out = subprocess.run([sys.executable, RECALC, case, '240'], capture_output=True, text=True).stdout
res = json.loads(out[out.index('{'):])
print('recalc:', res.get('status'), 'errors', res.get('total_errors'), 'formulas', res.get('total_formulas'))

# ---------------------------------------------------------------- the same arithmetic, independently
src = load_workbook(case)
s_st, s_de, s_pk = src['Stock'], src['Daily Entry'], src['Packing']
rows = []
for r in range(7, s_de.max_row + 1):
    d, item, q = (s_de.cell(row=r, column=c).value for c in (1, 2, 3))
    if isinstance(d, dt.datetime) and item and isinstance(q, (int, float)):
        rows.append((d, item, q))
last30 = max(COUNT, STMT - dt.timedelta(days=29))
window = max(1, min(30, (STMT - COUNT).days + 1))
prod = {}
for r in range(8, 102):
    a, b = s_st.cell(row=r, column=1).value, s_st.cell(row=r, column=2).value
    if not a: continue
    key = f'{a} - {b}'
    carton, per = s_pk.cell(row=r, column=4).value, s_pk.cell(row=r, column=5).value
    since = sum(q for d, i, q in rows if i == key and COUNT <= d <= STMT)
    today = sum(q for d, i, q in rows if i == key and d == STMT)
    l30 = sum(q for d, i, q in rows if i == key and last30 <= d <= STMT)
    ok = bool(carton) and isinstance(per, (int, float)) and per > 0
    used = math.ceil(since / per) if ok else None
    used_today = used - math.ceil((since - today) / per) if ok else None
    prod[r] = dict(key=key, carton=carton, per=per, since=since, today=today, used=used, used_today=used_today,
                   rate=(l30 / per if ok else 0), packed_no_carton=(since > 0 and not ok), unset=not ok)

fail = []
def eq(name, got, want, tol=1e-6):
    if want is None or want == '':
        if got not in (None, ''): fail.append(f'{name}: got {got!r} want blank')
    elif isinstance(want, str):
        if got != want: fail.append(f'{name}: got {got!r} want {want!r}')
    elif not isinstance(got, (int, float)) or abs(got - want) > tol:
        fail.append(f'{name}: got {got!r} want {want!r}')

v = load_workbook(case, data_only=True)
vcs, vpk = v['Carton Stock'], v['Packing']
for r, p in prod.items():
    eq(f'Packing F{r} {p["key"]}', vpk[f'F{r}'].value, p['since'])
    eq(f'Packing G{r}', vpk[f'G{r}'].value, p['used'])
    eq(f'Packing H{r}', vpk[f'H{r}'].value, p['today'])
    eq(f'Packing I{r}', vpk[f'I{r}'].value, p['used_today'])

names = {vcs.cell(row=r, column=2).value: r for r in range(8, 48) if vcs.cell(row=r, column=2).value}
for n, r in names.items():
    rec = sum(x or 0 for d, c, x, a in deliveries if c == n and COUNT <= d <= STMT)
    adj = sum(a or 0 for d, c, x, a in deliveries if c == n and COUNT <= d <= STMT)
    rec_today = sum(x or 0 for d, c, x, a in deliveries if c == n and d == STMT)
    mine = [p for p in prod.values() if p['carton'] == n and p['used'] is not None]
    used, used_today = sum(p['used'] for p in mine), sum(p['used_today'] for p in mine)
    o = opening.get(n)
    close = None if o is None else o + rec - used + adj
    lvl = reorder.get(n)
    status = ('Count needed' if o is None else 'Negative — check' if close < 0
              else 'Reorder' if lvl is not None and close <= lvl else 'OK')
    avg = sum(p['rate'] for p in mine) / window
    days = None if close is None else 0 if close <= 0 else ('—' if avg <= 0 else math.floor(close / avg))
    for col, want in zip('DEFGHIK', (rec_today, used_today, rec, used, adj, close, status)):
        eq(f'{n} {col}', vcs[f'{col}{r}'].value, want)
    eq(f'{n} days left', vcs[f'L{r}'].value, days)
    if o is not None:
        print(f'  {n:22} open {o:>5}  +recd {rec:>4}  −used {used:>5}  ±adj {adj:>3}  = {close:>6}   {status:<16} days left {days}')

warn = [vcs.cell(row=r, column=1).value or '' for r in range(50, 56)]
packed_no = sum(p['packed_no_carton'] for p in prod.values())
unset = sum(p['unset'] for p in prod.values())
eq('warning: packed with no carton', warn[0], f'{packed_no} product(s) packed since the count have no carton or no Packed per Carton on Packing — their cartons are not taken off. Fill them in on Packing.' if packed_no else '')
eq('warning: unknown carton on Packing', warn[1], '')
eq('warning: unknown carton on entry', warn[2], '1 row(s) on Carton Entry name a carton that is not on this list — they are not counted.')
eq('warning: duplicates', warn[3], '')
eq('warning: dates', warn[4], '')
eq('note: no carton size', warn[5], f'{unset} product(s) on Packing have no carton size yet — fill them in before they are packed.')
eq('entry check J4', v['Carton Entry']['J4'].value, 1)
eq('total closing', vcs['I48'].value, sum(vcs[f'I{r}'].value or 0 for r in range(8, 48)))
print(f'  products packed since the count with no carton set: {packed_no}; with no carton size at all: {unset}')
# ---------------------------------------------------------------- the live sheets are untouched
a, b = load_workbook(SRC), load_workbook(built)
sig = lambda c: (c.value, c.number_format, c.font.name, c.font.sz, c.font.b, str(c.font.color.rgb) if c.font.color else None,
                 c.fill.patternType, str(c.fill.fgColor.rgb), c.alignment.horizontal, c.protection.locked,
                 c.border.left.style if c.border.left else None)
for name in a.sheetnames:
    x, y = a[name], b[name]
    stop = 44 if name == 'Read me' else x.max_row          # the carton notes are appended below row 44
    bad = [c.coordinate for row in x.iter_rows(max_row=stop) for c in row if sig(c) != sig(y[c.coordinate])]
    if bad: fail.append(f'{name}: {len(bad)} cells changed, e.g. {bad[:3]}')
    for attr in ('freeze_panes', 'print_title_rows'):
        if getattr(x, attr) != getattr(y, attr): fail.append(f'{name}: {attr} changed')
    cf = lambda ws: sorted((str(k.sqref), tuple(tuple(r.formula) for r in v)) for k, v in ws.conditional_formatting._cf_rules.items())
    dvs = lambda ws: sorted((str(d.sqref), d.type, d.formula1) for d in ws.data_validations.dataValidation)
    if cf(x) != cf(y) or dvs(x) != dvs(y): fail.append(f'{name}: rules or dropdowns changed')
    if (x.protection.sheet, x.protection.formatCells, x.protection.password) != (y.protection.sheet, y.protection.formatCells, y.protection.password):
        fail.append(f'{name}: protection changed')
if not (b.security and b.security.lockStructure): fail.append('workbook structure lock lost')
if not b['Daily Report']._images: fail.append('Daily Report logo lost')
print(f'  live sheets unchanged: {", ".join(a.sheetnames)}')

print('\n'.join(fail) if fail else f'all checks passed — {len(prod)} products and {len(names)} cartons match the independent '
      f'calculation (count {COUNT:%d/%m/%Y} to statement {STMT:%d/%m/%Y}, {window} days averaged)')
sys.exit(1 if fail else 0)
